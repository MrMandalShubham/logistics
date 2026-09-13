# Builds the client deck. Run: python scripts/make-deck.mjs.py
#
# Deliberately light on technology and heavy on the story: what was
# wrong, what we built, what happens to one order, and what is true
# today. Nothing here claims more than the system has actually done.

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = r"C:\Users\shubh\Desktop\Logistics-Core-Overview.pptx"

INK    = RGBColor(0x10, 0x18, 0x20)
BLUE   = RGBColor(0x0B, 0x5F, 0xFF)
SKY    = RGBColor(0xE8, 0xF0, 0xFF)
GREY   = RGBColor(0x6B, 0x72, 0x80)
LIGHT  = RGBColor(0xF4, 0xF5, 0xF7)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
AMBER  = RGBColor(0xB3, 0x59, 0x00)
AMBERL = RGBColor(0xFF, 0xF4, 0xD6)
GREEN  = RGBColor(0x1A, 0x7F, 0x37)
GREENL = RGBColor(0xE8, 0xF5, 0xEC)
RED    = RGBColor(0x8A, 0x1C, 0x10)
REDL   = RGBColor(0xFF, 0xF0, 0xEE)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
W, H = 13.333, 7.5


def blank():
    return prs.slides.add_slide(prs.slide_layouts[6])


def rect(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=1.25):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None):
    """runs: list of (string, size, bold, colour) or (string, size, bold, colour, space_after)"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, r in enumerate(runs):
        body, size, bold, colour = r[0], r[1], r[2], r[3]
        after = r[4] if len(r) > 4 else 6
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(after)
        if spacing:
            p.line_spacing = spacing
        run = p.add_run(); run.text = body
        run.font.size = Pt(size); run.font.bold = bold
        run.font.color.rgb = colour; run.font.name = "Segoe UI"
    return tb


def fill_bg(s, colour):
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = colour
    bg.line.fill.background(); bg.shadow.inherit = False
    s.shapes._spTree.remove(bg._element)
    s.shapes._spTree.insert(2, bg._element)


def header(s, kicker, title, sub=None):
    text(s, 0.85, 0.62, 11.5, 0.4, [(kicker.upper(), 12, True, BLUE)])
    text(s, 0.85, 1.02, 11.5, 0.9, [(title, 34, True, INK)])
    if sub:
        text(s, 0.85, 1.92, 11.0, 0.6, [(sub, 15, False, GREY)])
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(1.85),
                            Inches(1.1), Inches(0.045))
    ln.fill.solid(); ln.fill.fore_color.rgb = BLUE
    ln.line.fill.background(); ln.shadow.inherit = False
    return s


def chip(s, x, y, w, h, label, body, fill, edge, label_colour):
    rect(s, x, y, w, h, fill, edge)
    text(s, x + 0.28, y + 0.22, w - 0.56, 0.35, [(label, 12, True, label_colour)])
    text(s, x + 0.28, y + 0.66, w - 0.56, h - 0.9, [(body, 12.5, False, INK)], spacing=1.15)


def arrow(s, x, y, w=0.55):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y),
                           Inches(w), Inches(0.26))
    a.fill.solid(); a.fill.fore_color.rgb = RGBColor(0xC7, 0xD2, 0xE0)
    a.line.fill.background(); a.shadow.inherit = False


def stat(s, x, y, w, big, cap, colour=BLUE):
    text(s, x, y, w, 1.0, [(big, 46, True, colour)], align=PP_ALIGN.CENTER)
    text(s, x, y + 0.95, w, 0.7, [(cap, 12, False, GREY)], align=PP_ALIGN.CENTER)


def footer(s, n):
    text(s, 11.9, 6.92, 0.9, 0.3, [(str(n), 10, False, RGBColor(0xC0, 0xC6, 0xCE))],
         align=PP_ALIGN.RIGHT)


# ══════════════════ 1. title ══════════════════
s = blank(); fill_bg(s, INK)
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.5),
                         Inches(1.5), Inches(0.06))
bar.fill.solid(); bar.fill.fore_color.rgb = BLUE
bar.line.fill.background(); bar.shadow.inherit = False

text(s, 0.9, 2.85, 11, 1.3, [("Logistics Core", 56, True, WHITE)])
text(s, 0.9, 4.15, 10, 0.9,
     [("The delivery system that connects your shop to your stockroom", 20, False,
       RGBColor(0x9A, 0xA5, 0xB4))])
text(s, 0.9, 5.6, 10, 0.5,
     [("Built, tested and ready to deploy", 14, False, RGBColor(0x6B, 0x77, 0x88))])

# ══════════════════ 2. the three systems ══════════════════
s = blank(); footer(s, 2)
header(s, "Where we started", "You had two systems that never spoke",
       "A customer could buy something. Nobody could tell you where it went.")

chip(s, 1.0, 2.9, 3.3, 2.0, "SHOP",
     "Customers browse, add to basket and pay. It reserves the stock it needs.",
     SKY, RGBColor(0xC7, 0xDA, 0xFF), BLUE)
chip(s, 9.0, 2.9, 3.3, 2.0, "STOCKROOM",
     "Knows what is on every shelf, in every shop, at every moment.",
     SKY, RGBColor(0xC7, 0xDA, 0xFF), BLUE)

gap = rect(s, 5.0, 2.9, 3.3, 2.0, REDL, RGBColor(0xF0, 0xB4, 0xAC))
text(s, 5.28, 3.35, 2.75, 1.2,
     [("?", 54, True, RED), ("nothing here", 13, False, RED)], align=PP_ALIGN.CENTER)

arrow(s, 4.45, 3.78); arrow(s, 8.42, 3.78)

text(s, 1.0, 5.35, 11.3, 1.2,
     [("Between paying and arriving, there was no system at all — no delivery record, "
       "no rider, no proof, and no way for your stockroom to learn that goods had left "
       "the building.", 15, False, INK)], spacing=1.3)

# ══════════════════ 3. the finding ══════════════════
s = blank(); fill_bg(s, INK); footer(s, 3)
text(s, 0.85, 0.62, 11.5, 0.4, [("WHAT WE FOUND FIRST", 12, True, BLUE)])
text(s, 0.85, 1.15, 11.5, 1.0, [("Every sale was quietly undoing itself", 34, True, WHITE)])

rect(s, 0.85, 2.5, 11.6, 2.1, RGBColor(0x1A, 0x24, 0x30), None)
text(s, 1.35, 2.75, 10.6, 1.7,
     [("When a customer paid, the shop set their goods aside in the stockroom.",
       17, False, RGBColor(0xD5, 0xDC, 0xE5)),
      ("Thirty minutes later, the stockroom put them back on the shelf — because "
       "nobody had ever told it the sale went through.", 17, True, WHITE)], spacing=1.3)

stat(s, 1.0, 5.0, 3.4, "30 min", "and the stock came back", RGBColor(0xFF, 0x9A, 0x3D))
stat(s, 5.0, 5.0, 3.4, "0", "sales ever recorded", RGBColor(0xFF, 0x9A, 0x3D))
stat(s, 9.0, 5.0, 3.4, "every", "order affected", RGBColor(0xFF, 0x9A, 0x3D))

# ══════════════════ 4. what we built ══════════════════
s = blank(); footer(s, 4)
header(s, "What we built", "One system in the middle",
       "It takes the order, gets it delivered, and tells both sides what happened.")

chip(s, 0.85, 3.1, 3.0, 1.7, "SHOP",
     "Hands over a paid order, with the address and the shop it comes from.",
     LIGHT, RGBColor(0xE0, 0xE3, 0xE8), GREY)

core = rect(s, 4.6, 2.8, 4.1, 2.3, BLUE, None)
text(s, 4.85, 3.15, 3.6, 1.6,
     [("LOGISTICS CORE", 13, True, RGBColor(0xBF, 0xD4, 0xFF)),
      ("Dispatch · Riders\nProof · Returns\nReports", 17, True, WHITE)],
     align=PP_ALIGN.CENTER, spacing=1.25)

chip(s, 9.45, 3.1, 3.0, 1.7, "STOCKROOM",
     "Told when goods truly leave, and when they come back.",
     LIGHT, RGBColor(0xE0, 0xE3, 0xE8), GREY)

arrow(s, 3.95, 3.85); arrow(s, 8.82, 3.85)

text(s, 0.85, 5.5, 11.6, 1.0,
     [("Neither of your existing systems was rebuilt. Nothing was replaced. "
       "We added the missing piece and connected it at both ends.", 15, False, INK)],
     spacing=1.3)

# ══════════════════ 5. the journey ══════════════════
s = blank(); footer(s, 5)
header(s, "One order, end to end", "The journey of a delivery")

steps = [
    ("Order arrives", "Paid, with an address and a shop"),
    ("Dispatcher assigns", "System suggests the nearest rider"),
    ("Rider collects", "Picks it up from the shop"),
    ("On the way", "Customer sees it moving"),
    ("At the door", "Customer reads out their code"),
    ("Delivered", "Stockroom told, and verified"),
]
x = 0.72
for i, (t, b) in enumerate(steps):
    rect(s, x, 2.9, 1.75, 1.9, WHITE, RGBColor(0xDD, 0xE1, 0xE7))
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.68), Inches(2.62),
                              Inches(0.42), Inches(0.42))
    circ.fill.solid(); circ.fill.fore_color.rgb = BLUE
    circ.line.fill.background(); circ.shadow.inherit = False
    text(s, x + 0.68, 2.66, 0.42, 0.35, [(str(i + 1), 13, True, WHITE)],
         align=PP_ALIGN.CENTER)
    text(s, x + 0.16, 3.25, 1.45, 0.6, [(t, 13, True, INK)], align=PP_ALIGN.CENTER)
    text(s, x + 0.16, 3.92, 1.45, 0.8, [(b, 10.5, False, GREY)],
         align=PP_ALIGN.CENTER, spacing=1.15)
    if i < 5:
        arrow(s, x + 1.80, 3.72, 0.32)
    x += 2.05

text(s, 0.72, 5.45, 11.8, 1.1,
     [("At every step the customer's order page updates, the stockroom stays in step, "
       "and a permanent record is written that nobody can edit afterwards.",
       15, False, INK)], spacing=1.3)

# ══════════════════ 6. dispatcher ══════════════════
s = blank(); footer(s, 6)
header(s, "For your dispatcher", "One screen, one obvious next action")

rows = [
    ("What is waiting", "Every paid order that needs a rider, oldest first."),
    ("Who is free", "Riders on shift, how many jobs each is carrying."),
    ("Who to pick", "The system ranks them — nearest first — and says why."),
    ("You still decide", "It suggests. The dispatcher chooses and clicks."),
]
y = 2.8
for t, b in rows:
    rect(s, 0.85, y, 11.6, 0.92, LIGHT, None)
    text(s, 1.25, y + 0.17, 3.3, 0.5, [(t, 15, True, INK)])
    text(s, 4.8, y + 0.2, 7.3, 0.5, [(b, 13.5, False, GREY)])
    y += 1.05

text(s, 0.85, 6.55, 11.6, 0.6,
     [("\"0.6 km away, last seen just now, 0 of 5 jobs in hand\"", 14, True, BLUE)])

# ══════════════════ 7. rider ══════════════════
s = blank(); footer(s, 7)
header(s, "For your riders", "A phone app that works in a basement",
       "No app store. It opens in the phone's browser and behaves like an app.")

chip(s, 0.85, 2.95, 3.6, 2.35, "BIG BUTTONS",
     "One job, one obvious next action, sized for a thumb on a doorstep.",
     WHITE, RGBColor(0xDD, 0xE1, 0xE7), BLUE)
chip(s, 4.85, 2.95, 3.6, 2.35, "WORKS OFFLINE",
     "Every tap is saved on the phone first, then sent when signal returns. "
     "Nothing is lost in a lift or a car park.",
     GREENL, RGBColor(0xB7, 0xE0, 0xC4), GREEN)
chip(s, 8.85, 2.95, 3.6, 2.35, "HONEST STATUS",
     "It says \"waiting to sync\", never \"delivered\", until the system has "
     "actually agreed.",
     AMBERL, RGBColor(0xFF, 0xD6, 0x99), AMBER)

text(s, 0.85, 5.7, 11.6, 0.9,
     [("The time on the record is when the rider tapped the button — not when the "
       "phone found signal. Otherwise every delivery time you measure would be wrong.",
       14, False, INK)], spacing=1.3)

# ══════════════════ 8. the code ══════════════════
s = blank(); footer(s, 8)
header(s, "Proof of delivery", "A six-digit code at the door")

rect(s, 0.85, 2.9, 5.5, 2.9, WHITE, RGBColor(0xDD, 0xE1, 0xE7))
text(s, 1.25, 3.2, 4.7, 2.4,
     [("How it works", 15, True, INK),
      ("The customer gets a code when the rider arrives.", 13.5, False, GREY),
      ("They read it out. The rider types it in.", 13.5, False, GREY),
      ("The system checks it before the delivery counts.", 13.5, False, GREY)],
     spacing=1.45)

rect(s, 6.95, 2.9, 5.5, 2.9, SKY, RGBColor(0xC7, 0xDA, 0xFF))
text(s, 7.35, 3.2, 4.7, 2.4,
     [("Why it is safe", 15, True, BLUE),
      ("The rider never sees the code in the app.", 13.5, False, INK),
      ("It expires, and it works only once.", 13.5, False, INK),
      ("It is never written down anywhere readable.", 13.5, False, INK)],
     spacing=1.45)

text(s, 0.85, 6.2, 11.6, 0.8,
     [("If the code is wrong, the delivery does not simply fail — it is flagged for a "
       "person to look at. The parcel is gone either way; whether the proof was good "
       "is not a decision for a machine.", 14, False, INK)], spacing=1.3)

# ══════════════════ 9. customer ══════════════════
s = blank(); footer(s, 9)
header(s, "For your customers", "Three updates, not eight",
       "Their order page finally tells them something true.")

msgs = [
    ("Being packed", "We have your order and it is being packed.", SKY, BLUE),
    ("On its way", "Your order is on its way.", SKY, BLUE),
    ("Delivered", "Delivered. Thank you.", GREENL, GREEN),
]
x = 0.85
for t, m, bg, c in msgs:
    rect(s, x, 2.9, 3.7, 1.7, bg, None)
    text(s, x + 0.3, 3.12, 3.1, 0.4, [(t.upper(), 11, True, c)])
    text(s, x + 0.3, 3.55, 3.1, 0.9, [("\u201c" + m + "\u201d", 14, False, INK)],
         spacing=1.25)
    x += 3.9

rect(s, 0.85, 5.0, 11.6, 1.25, AMBERL, RGBColor(0xFF, 0xD6, 0x99))
text(s, 1.25, 5.25, 10.8, 0.9,
     [("And when it does not arrive, they are told that too — with the real reason. "
       "\u201cWe could not reach you at the address. We will try again.\u201d",
       14, False, AMBER)], spacing=1.25)

text(s, 0.85, 6.5, 11.6, 0.6,
     [("Before this, every order that was not awaiting payment displayed as "
       "\u201cDelivered\u201d — including cancelled ones.", 13, False, GREY)])

# ══════════════════ 10. when it goes wrong ══════════════════
s = blank(); footer(s, 10)
header(s, "When it goes wrong", "Somebody decides. Never a timer.")

flow = [
    ("Nobody home", REDL, RED),
    ("Flagged to a person", AMBERL, AMBER),
    ("Try again, or bring it back", SKY, BLUE),
    ("Stockroom told", GREENL, GREEN),
]
x = 0.85
for t, bg, c in flow:
    rect(s, x, 3.0, 2.6, 1.5, bg, None)
    text(s, x + 0.25, 3.35, 2.1, 0.9, [(t, 14, True, c)],
         align=PP_ALIGN.CENTER, spacing=1.2)
    if t != "Stockroom told":
        arrow(s, x + 2.66, 3.62, 0.42)
    x += 3.0

text(s, 0.85, 5.1, 11.6, 1.6,
     [("Every decision is recorded with a name and a reason.", 16, True, INK),
      ("When two people disagree about what happened at a door — a rider says they "
       "delivered it, the system says somebody else had the parcel — nothing is "
       "quietly overwritten. Both accounts are kept and a person decides.",
       14, False, GREY)], spacing=1.3)

# ══════════════════ 11. the money ══════════════════
s = blank(); footer(s, 11)
header(s, "The part that pays for itself", "Your stockroom finally knows the truth")

rect(s, 0.85, 2.85, 5.55, 2.55, REDL, RGBColor(0xF0, 0xB4, 0xAC))
text(s, 1.25, 3.1, 4.75, 2.1,
     [("BEFORE", 12, True, RED),
      ("Goods left the building and the stockroom never knew.", 14, False, INK),
      ("Stock counts drifted from reality.", 14, False, INK),
      ("Returns were invisible.", 14, False, INK)], spacing=1.4)

rect(s, 6.9, 2.85, 5.55, 2.55, GREENL, RGBColor(0xB7, 0xE0, 0xC4))
text(s, 7.3, 3.1, 4.75, 2.1,
     [("NOW", 12, True, GREEN),
      ("A delivery is recorded as a sale — and checked afterwards.", 14, False, INK),
      ("A return puts the goods back, and that is checked too.", 14, False, INK),
      ("If the two ever disagree, somebody is told.", 14, False, INK)], spacing=1.4)

text(s, 0.85, 5.75, 11.6, 1.0,
     [("We do not trust a \u201cyes\u201d from the stockroom. Every time we record a "
       "sale, we read it back and confirm the goods really moved. A mismatch becomes "
       "somebody's job, not a silent error.", 14, False, INK)], spacing=1.3)

# ══════════════════ 12. principles ══════════════════
s = blank(); fill_bg(s, INK); footer(s, 12)
text(s, 0.85, 0.62, 11.5, 0.4, [("HOW IT IS BUILT", 12, True, BLUE)])
text(s, 0.85, 1.15, 11.5, 0.9, [("Four rules, everywhere", 34, True, WHITE)])

rules = [
    ("Nothing is ever deleted", "The history of a delivery cannot be edited or removed — by anyone."),
    ("Nothing is silently lost", "A message that fails is retried, then escalated to a person."),
    ("Never guess", "If the system does not know something, it says so instead of inventing it."),
    ("A person decides", "Anything ambiguous goes to a human with the full picture."),
]
y = 2.45
for t, b in rules:
    rect(s, 0.85, y, 11.6, 1.0, RGBColor(0x1A, 0x24, 0x30), None)
    text(s, 1.3, y + 0.2, 4.0, 0.5, [(t, 16, True, WHITE)])
    text(s, 5.5, y + 0.24, 6.6, 0.6, [(b, 13, False, RGBColor(0x9A, 0xA5, 0xB4))])
    y += 1.15

# ══════════════════ 13. confidence ══════════════════
s = blank(); footer(s, 13)
header(s, "How we know it works", "Checked, not hoped")

stat(s, 0.85, 2.9, 3.6, "345", "automated checks, run on every change")
stat(s, 4.85, 2.9, 3.6, "8", "stages, each reviewed and signed off")
stat(s, 8.85, 2.9, 3.6, "0", "known defects outstanding")

rect(s, 0.85, 5.0, 11.6, 1.65, LIGHT, None)
text(s, 1.3, 5.25, 10.8, 1.25,
     [("Every claim in this deck was proven against a running system, not a diagram.",
       15, True, INK),
      ("A real delivery reduced real stock. A real return put it back. A rider "
       "completed a job with no signal and lost nothing.", 13.5, False, GREY)],
     spacing=1.3)

# ══════════════════ 14. what it tells you ══════════════════
s = blank(); footer(s, 14)
header(s, "What it tells you", "The questions you can now answer")

qs = [
    "Where is every order, right now?",
    "What has been sitting still for an hour?",
    "How long do we take, from paying to doorstep?",
    "Which reasons cause the most failed deliveries?",
    "Did every delivery reach the stockroom correctly?",
    "Which areas are we failing in most often?",
]
x, y = 0.85, 2.9
for i, q in enumerate(qs):
    rect(s, x, y, 5.55, 0.95, WHITE, RGBColor(0xDD, 0xE1, 0xE7))
    text(s, x + 0.35, y + 0.25, 5.0, 0.5, [(q, 13.5, False, INK)])
    if i % 2 == 0:
        x += 6.05
    else:
        x = 0.85; y += 1.1

text(s, 0.85, 6.35, 11.6, 0.6,
     [("Where a number cannot be trusted, the report says so rather than showing a "
       "figure somebody might act on.", 13, False, GREY)])

# ══════════════════ 15. status ══════════════════
s = blank(); footer(s, 15)
header(s, "Where it stands today", "Built and tested. Not yet live.")

ready = [
    "The delivery system is complete and connected at both ends",
    "The shop's changes are merged and deployed",
    "The database is installed and verified",
]
todo = [
    "Choose where to host it",
    "Enter each shop's location on a map",
    "Run one real order end to end before opening it up",
]

rect(s, 0.85, 2.85, 5.55, 3.0, GREENL, RGBColor(0xB7, 0xE0, 0xC4))
text(s, 1.25, 3.1, 4.75, 0.4, [("DONE", 12, True, GREEN)])
y = 3.6
for r in ready:
    text(s, 1.25, y, 4.75, 0.7, [("\u2713   " + r, 13, False, INK)], spacing=1.2)
    y += 0.72

rect(s, 6.9, 2.85, 5.55, 3.0, AMBERL, RGBColor(0xFF, 0xD6, 0x99))
text(s, 7.3, 3.1, 4.75, 0.4, [("BEFORE GOING LIVE", 12, True, AMBER)])
y = 3.6
for r in todo:
    text(s, 7.3, y, 4.75, 0.7, [("\u25A1   " + r, 13, False, INK)], spacing=1.2)
    y += 0.72

text(s, 0.85, 6.2, 11.6, 0.8,
     [("None of what remains is development work. It is deployment and setup — "
       "measured in days, not weeks.", 14, True, INK)])

# ══════════════════ 16. close ══════════════════
s = blank(); fill_bg(s, INK)
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.6),
                         Inches(1.5), Inches(0.06))
bar.fill.solid(); bar.fill.fore_color.rgb = BLUE
bar.line.fill.background(); bar.shadow.inherit = False

text(s, 0.9, 2.95, 11.2, 2.2,
     [("You can now tell a customer\nwhere their order is —", 40, True, WHITE),
      ("and your stockroom finally agrees.", 40, True, BLUE)], spacing=1.15)

text(s, 0.9, 5.85, 10, 0.5,
     [("Logistics Core", 14, False, RGBColor(0x6B, 0x77, 0x88))])

prs.save(OUT)
print("written:", OUT)
print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))
